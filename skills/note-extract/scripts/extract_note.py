#!/usr/bin/env python3
"""Extract XHS notes without cookies first; authenticated retry exports MD/DOCX."""
import argparse
import hashlib
import html
import io
import json
import os
import re
import shutil
import sqlite3
import ssl
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime
from html.parser import HTMLParser
from http.cookiejar import Cookie, CookieJar
from pathlib import Path
from zoneinfo import ZoneInfo

SKILL = Path(__file__).resolve().parents[1]
dep = SKILL / '.deps'
if not dep.is_dir():
    dep = SKILL.parents[1] / '.runtime/note-extract-deps'  # Existing project installs.
if dep.is_dir():
    sys.path.insert(0, str(dep))

TZ = ZoneInfo('Asia/Shanghai')
UA = 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/140.0.0.0 Safari/537.36'
PAGE_HEADERS = {
    'User-Agent': UA,
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8',
    'Accept-Language': 'zh-CN,zh;q=0.9,en;q=0.8',
    'Referer': 'https://www.xiaohongshu.com/',
    'Upgrade-Insecure-Requests': '1',
}


class ExtractError(Exception):
    def __init__(self, code, message, http_status=None):
        self.code = code
        self.http_status = http_status
        super().__init__(message)


def emit(**data):
    print(json.dumps(data, ensure_ascii=False), flush=True)


def allowed_host(host, base):
    return host == base or host.endswith('.' + base)


def clean_url(value):
    value = html.unescape(value.strip()).replace('\\:', ':').replace('\\.', '.').replace('\\&', '&')
    found = re.search(r'https?://[^\s<>"\\]+', value)
    if not found:
        raise ExtractError('INVALID_URL', '没有找到小红书链接。')
    value = found.group().rstrip('，。')
    p = urllib.parse.urlsplit(value)
    if p.username or p.password or p.port not in (None, 443):
        raise ExtractError('INVALID_URL', '链接主机或端口不正确。')
    if not any(allowed_host(p.hostname or '', h) for h in ('xiaohongshu.com', 'xhslink.com')):
        raise ExtractError('INVALID_URL', '仅支持小红书笔记链接。')
    return urllib.parse.urlunsplit(('https', p.netloc, p.path, p.query, ''))


class RestrictedRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        p = urllib.parse.urlsplit(newurl)
        if p.scheme != 'https' or not any(allowed_host(p.hostname or '', h) for h in ('xiaohongshu.com', 'xhslink.com', 'xhscdn.com')):
            raise ExtractError('UNEXPECTED_REDIRECT', '请求重定向到非预期站点，已停止。')
        redirected = super().redirect_request(req, fp, code, msg, headers, newurl)
        if redirected:
            redirected.remove_header('Cookie')
        return redirected


def anonymous_fetch_curl(url, image=False):
    """Use the system HTTP client without its config, cookies, or unchecked redirects."""
    limit = 40 * 1024 * 1024 if image else 12 * 1024 * 1024
    headers = dict(PAGE_HEADERS)
    if image:
        headers['Accept'] = 'image/webp,image/png,image/jpeg,*/*;q=0.8'
    current = url
    with tempfile.TemporaryDirectory(prefix='note-extract-http-') as tmp:
        body, response_headers = Path(tmp) / 'body', Path(tmp) / 'headers'
        for _ in range(6):
            parsed = urllib.parse.urlsplit(current)
            if parsed.scheme != 'https' or parsed.username or parsed.password or parsed.port not in (None, 443) or not any(
                allowed_host(parsed.hostname or '', h) for h in ('xiaohongshu.com', 'xhslink.com', 'xhscdn.com')
            ):
                raise ExtractError('UNEXPECTED_REDIRECT', '请求目标不是预期的HTTPS站点，已停止。')
            command = ['curl', '-q', '--silent', '--show-error', '--max-time', '35',
                       '--proto', '=https', '--max-filesize', str(limit), '--header', 'Cookie:',
                       '--output', str(body), '--dump-header', str(response_headers),
                       '--write-out', '%{http_code}\n%{content_type}']
            for name, value in headers.items():
                command += ['--header', f'{name}: {value}']
            command.append(current)
            result = subprocess.run(command, capture_output=True, text=True)
            if result.returncode == 60:
                raise ExtractError('TLS_CERTIFICATE_ERROR', 'TLS证书验证失败，请检查本机CA配置；这不是登录失败。')
            if result.returncode == 63:
                raise ExtractError('RESPONSE_TOO_LARGE', '响应超过文件大小上限。')
            if result.returncode:
                raise ExtractError('NETWORK_ERROR', '网络请求失败或超时。')
            status_line, _, content_type = result.stdout.partition('\n')
            status = int(status_line)
            if status in (301, 302, 303, 307, 308):
                locations = re.findall(r'^location:\s*(.+)$', response_headers.read_text(errors='replace'), re.I | re.M)
                if not locations:
                    raise ExtractError('UNEXPECTED_REDIRECT', '重定向缺少目标地址，已停止。')
                current = urllib.parse.urljoin(current, locations[-1].strip())
                continue  # Next target is checked before another request is sent.
            if status >= 400:
                raise ExtractError('HTTP_ERROR', f'HTTP {status}，请检查笔记可访问性或刷新分享链接。', http_status=status)
            if body.stat().st_size > limit:
                raise ExtractError('RESPONSE_TOO_LARGE', '响应超过文件大小上限。')
            return body.read_bytes(), current, status, content_type.strip()
    raise ExtractError('UNEXPECTED_REDIRECT', '重定向次数超过上限，已停止。')


def fetch(url, jar=None, image=False):
    if jar is None and shutil.which('curl'):
        return anonymous_fetch_curl(url, image)
    context = ssl.create_default_context()
    try:
        import certifi
        context.load_verify_locations(cafile=certifi.where())
    except ImportError:
        pass  # System trust remains enabled when optional CA roots are unavailable.
    handlers = [RestrictedRedirect(), urllib.request.HTTPSHandler(context=context)]
    if jar is not None:
        handlers.append(urllib.request.HTTPCookieProcessor(jar))
    headers = dict(PAGE_HEADERS)
    if image:
        headers['Accept'] = 'image/webp,image/png,image/jpeg,*/*;q=0.8'
    try:
        with urllib.request.build_opener(*handlers).open(urllib.request.Request(url, headers=headers), timeout=35) as response:
            limit = 40 * 1024 * 1024 if image else 12 * 1024 * 1024
            data = response.read(limit + 1)
            if len(data) > limit:
                raise ExtractError('RESPONSE_TOO_LARGE', '响应超过文件大小上限。')
            return data, response.url, response.status, response.headers.get('Content-Type', '')
    except urllib.error.HTTPError as error:
        raise ExtractError('HTTP_ERROR', f'HTTP {error.code}，请检查笔记可访问性或刷新分享链接。', http_status=error.code) from None
    except urllib.error.URLError as error:
        if isinstance(error.reason, ssl.SSLCertVerificationError):
            raise ExtractError('TLS_CERTIFICATE_ERROR', 'TLS证书验证失败，请检查本机CA配置或安装certifi；这不是登录失败。') from None
        raise ExtractError('NETWORK_ERROR', '网络请求失败或超时。') from None
    except TimeoutError:
        raise ExtractError('NETWORK_ERROR', '网络请求失败或超时。') from None


def browser_cookie_jar(browser='chrome', profile=None):
    """Each profile is independent; never merge multiple account sessions."""
    candidates = [None]
    if sys.platform == 'darwin' and browser in ('chrome', 'edge', 'brave'):
        folders = {'chrome': 'Google/Chrome', 'edge': 'Microsoft Edge', 'brave': 'BraveSoftware/Brave-Browser'}
        root = Path.home() / 'Library/Application Support' / folders[browser]
        try:
            # This preflight distinguishes OS privacy denial from a logged-out profile.
            state = root / 'Local State'
            cache = json.loads(state.read_text()).get('profile', {}) if state.exists() else {}
            names = [profile] if profile else list(dict.fromkeys([cache.get('last_used', 'Default'), 'Default', *cache.get('info_cache', {})]))
            if any('/' in n or '\\' in n or n in ('.', '..') for n in names):
                raise ExtractError('INVALID_PROFILE', 'Chrome 配置名称不正确。')
            candidates = []
            for name in names:
                for suffix in ('Network/Cookies', 'Cookies'):
                    db = root / name / suffix
                    if db.exists():
                        with sqlite3.connect(db.as_uri() + '?mode=ro', uri=True, timeout=3) as con:
                            con.execute("SELECT name FROM cookies WHERE host_key IN ('xiaohongshu.com','.xiaohongshu.com','www.xiaohongshu.com','.www.xiaohongshu.com') LIMIT 1").fetchall()
                        candidates.append(db)
        except PermissionError:
            raise ExtractError('BROWSER_PERMISSION_DENIED', f'系统拒绝读取 {browser} 数据目录；这不代表未登录。请允许读取后重试，或提供本次请求的 Cookie。') from None
        except sqlite3.OperationalError:
            raise ExtractError('BROWSER_DB_UNAVAILABLE', '浏览器 Cookie 数据库被占用或无法读取；请处理后重试。') from None
    if not candidates:
        raise ExtractError('LOGIN_REQUIRED', f'未找到 {browser} 的登录 Cookie。请先在你自己的 {browser} 打开小红书并登录，登录后再继续提取。')
    try:
        import browser_cookie3
    except ImportError:
        raise ExtractError('DEPENDENCY_MISSING', '缺少 browser-cookie3，请安装 scripts/requirements.txt 到技能 .deps。') from None
    for candidate in candidates:
        try:
            jar = getattr(browser_cookie3, browser)(cookie_file=str(candidate) if candidate else None, domain_name='xiaohongshu.com')
        except PermissionError:
            raise ExtractError('BROWSER_PERMISSION_DENIED', f'系统拒绝读取 {browser} Cookie。') from None
        except Exception:
            # Never expose library exceptions containing decrypted values or key material.
            raise ExtractError('BROWSER_DECRYPT_FAILED', f'无法读取或解密 {browser} Cookie；如有系统授权弹窗请由用户处理。') from None
        filtered = CookieJar()
        for cookie in jar:
            domain = cookie.domain.lstrip('.')
            if domain in ('xiaohongshu.com', 'www.xiaohongshu.com') and not cookie.is_expired():
                filtered.set_cookie(cookie)
        if any(c.name == 'web_session' and c.value for c in filtered):
            return filtered, {'source': browser, 'profile': candidate.parent.name if candidate and candidate.parent.name != 'Network' else candidate.parent.parent.name if candidate else 'default', 'cookie_count': len(filtered)}
    raise ExtractError('LOGIN_REQUIRED', f'没有有效的小红书登录 Cookie。请先在你自己的 {browser} 打开小红书并登录，登录后告诉我再继续。')


def supplied_cookie_jar(raw):
    raw = raw.strip()
    if '\n' in raw or '\r' in raw:
        raise ExtractError('INVALID_COOKIE', 'Cookie 应为单行请求 Cookie 标头。')
    jar = CookieJar()
    for part in raw.split(';'):
        name, sep, value = part.strip().partition('=')
        if sep and re.fullmatch(r'[!#$%&\x27*+.^_`|~0-9A-Za-z-]+', name):
            jar.set_cookie(Cookie(0, name, value, None, False, '.xiaohongshu.com', True, True, '/', True, True, None, True, None, None, {}, False))
    if not any(c.name == 'web_session' and c.value for c in jar):
        raise ExtractError('LOGIN_REQUIRED', '提供的 Cookie 中没有小红书 web_session；请先在自己的浏览器登录。')
    return jar


class ScriptParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=False)
        self.scripts, self.current = [], None

    def handle_starttag(self, tag, attrs):
        if tag == 'script':
            self.current = [dict(attrs), '']

    def handle_data(self, data):
        if self.current is not None:
            self.current[1] += data

    def handle_endtag(self, tag):
        if tag == 'script' and self.current is not None:
            self.scripts.append(self.current)
            self.current = None


def parse_json_js(value):
    value = value.strip().rstrip(';')
    value = re.sub(r'"(?:\\.|[^"\\])*"|\bundefined\b', lambda m: 'null' if m[0] == 'undefined' else m[0], value)
    return json.loads(value)


def note_id(url):
    m = re.search(r'/(?:explore|discovery/item)/([0-9a-f]{24})(?:/|$)', urllib.parse.urlsplit(url).path)
    if not m:
        raise ExtractError('INVALID_NOTE_URL', '链接没有可识别的笔记 ID，请发送完整笔记分享链接。')
    return m[1]


def parse_page(page, url):
    target = note_id(url)
    parser = ScriptParser()
    parser.feed(page)
    article, note = None, None
    for attrs, text in parser.scripts:
        if attrs.get('type') == 'application/ld+json':
            try:
                data = json.loads(html.unescape(text))
                objects = data if isinstance(data, list) else data.get('@graph', [data])
                for obj in objects:
                    if obj.get('@type') == 'Article':
                        entity = obj.get('mainEntityOfPage', {})
                        identity = entity.get('@id', '') if isinstance(entity, dict) else str(entity)
                        if target in identity:
                            article = obj
            except (ValueError, TypeError, AttributeError):
                pass
        marker = re.search(r'(?:window\.)?__INITIAL_STATE__\s*=\s*', text)
        if marker:
            # Other stores can contain `new Map([])`; decode only noteDetailMap,
            # rather than executing JS or requiring the entire state to be JSON.
            state_text = re.sub(r'"(?:\\.|[^"\\])*"|\bundefined\b', lambda m: 'null' if m[0] == 'undefined' else m[0], text[marker.end():])
            for location in re.finditer(r'"noteDetailMap"\s*:\s*', state_text):
                try:
                    mapping, _ = json.JSONDecoder().raw_decode(state_text, location.end())
                    candidate = mapping.get(target, {}).get('note')
                    if candidate and candidate.get('noteId') == target:
                        note = candidate
                        break
                except (ValueError, TypeError, AttributeError):
                    pass
    if note:
        if note.get('type') not in ('normal', None, ''):
            raise ExtractError('VIDEO_NOT_SUPPORTED', '这是一篇视频笔记，当前技能只完整归档图文笔记。')
        raw_images = note.get('imageList', [])
        result = {
            'note_id': target, 'title': note.get('title', ''), 'raw_description': note.get('desc', ''),
            'author': note.get('user', {}).get('nickname', ''), 'author_id': note.get('user', {}).get('userId', ''),
            'published_at': timestamp(note.get('time')), 'modified_at': timestamp(note.get('lastUpdateTime')),
            'published_at_source': 'note.time', 'extraction_method': 'http-initial-state',
            'topics': [t['name'] for t in note.get('tagList', []) if t.get('name')],
        }
    elif article:
        raw_images = article.get('image', [])
        raw_images = raw_images if isinstance(raw_images, list) else [raw_images]
        author = article.get('author', {})
        if isinstance(author, list):
            author = author[0] if author else {}
        result = {
            'note_id': target, 'title': re.sub(r'\s*-\s*小红书$', '', article.get('headline', '')),
            'raw_description': article.get('description', ''),
            'author': author.get('name', '') if isinstance(author, dict) else str(author),
            'author_id': '', 'published_at': '', 'modified_at': '', 'published_at_source': 'unverified',
            'extraction_method': 'http-json-ld', 'topics': [],
        }
    else:
        raise ExtractError('NOTE_NOT_FOUND', 'HTTP 页面没有目标笔记正文，可能是首页壳页、登录失效或笔记不可访问；不能导出成功文档。')
    images, seen = [], set()
    for item in raw_images:
        if isinstance(item, str):
            item = {'urlDefault': item}
        src = item.get('urlDefault') or item.get('url') or item.get('contentUrl')
        if not src:
            src = next((v.get('url') for v in item.get('infoList', []) if v.get('imageScene') == 'WB_DFT'), None)
        if not src:
            raise ExtractError('IMAGE_URL_MISSING', '笔记图片列表存在没有下载地址的图片。')
        src = html.unescape(src).replace('http://', 'https://', 1)
        parsed = urllib.parse.urlsplit(src)
        if parsed.scheme != 'https' or not allowed_host(parsed.hostname or '', 'xhscdn.com'):
            raise ExtractError('UNEXPECTED_IMAGE_HOST', '图片不在已支持的小红书 CDN 中，已停止该篇。')
        key = item.get('fileId') or parsed.path.split('!')[0].rsplit('/', 1)[-1]
        if key in seen:
            continue
        seen.add(key)
        images.append({'file_id': key, 'url': src, 'declared_width': item.get('width'), 'declared_height': item.get('height')})
    if not (result['title'] or result['raw_description']) or not images:
        raise ExtractError('INCOMPLETE_NOTE', '笔记标题、正文或图片列表不完整。')
    result.update({
        'schema_version': 1, 'source_url': url, 'canonical_url': f'https://www.xiaohongshu.com/explore/{target}',
        'extracted_at': datetime.now(TZ).isoformat(timespec='seconds'),
        'body': re.sub(r'#([^#\n]*?)\[话题\]#', r'#\1', result['raw_description']),
        'images': images, 'source_image_count': len(raw_images),
        'duplicate_image_entries_removed': len(raw_images) - len(images),
        'raw_ld_dates': {k: article.get(k) for k in ('datePublished', 'dateModified')} if article else {},
    })
    return result


def timestamp(value):
    if not isinstance(value, (int, float)) or value <= 0:
        return ''
    return datetime.fromtimestamp(value / 1000 if value > 1e11 else value, TZ).isoformat(timespec='seconds')


def download_images(note, directory):
    from PIL import Image
    originals = directory / 'images/original'
    originals.mkdir(parents=True)
    for index, item in enumerate(note['images'], 1):
        data, _, status, content_type = fetch(item['url'], image=True)
        if not content_type.lower().startswith('image/'):
            raise ExtractError('INVALID_IMAGE', f'第 {index} 张图片返回非图片内容。')
        try:
            with Image.open(io.BytesIO(data)) as check:
                check.verify()
            picture = Image.open(io.BytesIO(data))
            picture.load()
        except Exception:
            raise ExtractError('INVALID_IMAGE', f'第 {index} 张图片无法解码。') from None
        extension = {'JPEG': 'jpg', 'PNG': 'png', 'WEBP': 'webp', 'GIF': 'gif', 'AVIF': 'avif'}.get(picture.format)
        if not extension:
            raise ExtractError('UNSUPPORTED_IMAGE', f'第 {index} 张图片格式不支持。')
        original = originals / f'{index:02d}.{extension}'
        original.write_bytes(data)
        png = directory / f'images/{index:02d}.png'
        picture.convert('RGBA' if 'A' in picture.getbands() else 'RGB').save(png, 'PNG')
        item.update({'position': index, 'original_file': str(original.relative_to(directory)), 'file': str(png.relative_to(directory)),
                     'width': picture.width, 'height': picture.height, 'content_type': content_type,
                     'sha256': hashlib.sha256(data).hexdigest(), 'png_sha256': hashlib.sha256(png.read_bytes()).hexdigest(), 'http_status': status})


def md_escape(text):
    return re.sub(r'([\\`*_{}\[\]<>#+.!|~-])', r'\\\1', text)


def write_markdown(note, directory):
    lines = ['# ' + md_escape(note['title'] or '无标题笔记'), '', f"作者：{md_escape(note['author'])}  "]
    if note['published_at']:
        lines.append(f"发布时间：{note['published_at'].replace('T', ' ')}  ")
    lines += [f"原文：[查看小红书笔记]({note['source_url']})  ", f"笔记 ID：{note['note_id']}", '', '## 正文', '']
    lines += [md_escape(line) + ('  ' if line else '') for line in note['body'].split('\n')]
    lines += ['', '## 笔记图片', '']
    for item in note['images']:
        lines += [f"### 原图 {item['position']:02d}", '', f"![原图 {item['position']:02d}]({item['file']})", '']
    (directory / '笔记.md').write_text('\n'.join(lines), encoding='utf-8')


def write_docx(note, directory):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Arial Unicode MS'
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + attr), None)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(4)
    doc.styles['Title'].font.size = Pt(20)
    doc.styles['Heading 1'].font.size = Pt(13)
    doc.add_paragraph(note['title'] or '无标题笔记', 'Title')
    meta = doc.add_paragraph(f"作者：{note['author']}")
    meta.paragraph_format.space_after = Pt(2)
    if note['published_at']:
        doc.add_paragraph('发布时间：' + note['published_at'].replace('T', ' '))
    p = doc.add_paragraph('原文：')
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(note['source_url'], RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    txt = OxmlElement('w:t')
    txt.text = '查看小红书笔记'
    run.append(txt)
    link.append(run)
    p._p.append(link)
    doc.add_heading('正文', level=1)
    for line in note['body'].split('\n'):
        doc.add_paragraph(line)
    for item in note['images']:
        doc.add_page_break()
        doc.add_heading(f"原图 {item['position']:02d}", level=1)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        width = min(6.3, 8.4 * item['width'] / item['height'])
        picture = p.add_run().add_picture(str(directory / item['file']), width=Inches(width))
        picture._inline.docPr.set('descr', f"{note['title']} 原图 {item['position']}")
    doc.core_properties.title = note['title']
    doc.core_properties.author = note['author']
    doc.core_properties.subject = '小红书笔记原文与图片归档'
    doc.save(directory / '笔记.docx')


def verify(note, directory):
    from docx import Document
    doc = Document(directory / '笔记.docx')
    if len(doc.inline_shapes) != len(note['images']):
        raise ExtractError('VERIFY_FAILED', 'Word 嵌入图片数与笔记不一致。')
    texts = [p.text for p in doc.paragraphs]
    start = texts.index('正文') + 1
    body_lines = note['body'].split('\n')
    if texts[start:start + len(body_lines)] != body_lines:
        raise ExtractError('VERIFY_FAILED', 'Word 正文与原始笔记不一致。')
    for shape, item in zip(doc.inline_shapes, note['images']):
        rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
        actual = hashlib.sha256(doc.part.related_parts[rid].blob).hexdigest()
        if actual != item['png_sha256']:
            raise ExtractError('VERIFY_FAILED', 'Word 图片顺序或嵌入内容不一致。')
    md = (directory / '笔记.md').read_text()
    for item in note['images']:
        if f"]({item['file']})" not in md or not (directory / item['file']).is_file():
            raise ExtractError('VERIFY_FAILED', 'Markdown 图片引用缺失。')
    with zipfile.ZipFile(directory / '笔记.docx') as archive:
        if archive.testzip() is not None:
            raise ExtractError('VERIFY_FAILED', 'Word ZIP 完整性校验失败。')


def export_note(note, root):
    title = re.sub(r'[\x00-\x1f/\\:*?"<>|]', '_', note['title'] or '无标题笔记').strip('. ')[:65]
    directory = root / f"{title}_{note['note_id']}"
    if directory.exists():
        directory = root / f"{title}_{note['note_id']}_{datetime.now(TZ).strftime('%H%M%S_%f')}"
    directory.mkdir(parents=True)
    try:
        download_images(note, directory)
        write_markdown(note, directory)
        write_docx(note, directory)
        verify(note, directory)
        note['validation'] = {'content': 'passed', 'images': 'passed', 'docx_embedding': 'passed', 'visual_review': 'pending'}
        (directory / 'note.json').write_text(json.dumps(note, ensure_ascii=False, indent=2), encoding='utf-8')
        return directory
    except Exception:
        (directory / '提取未完成.txt').write_text('本篇提取未完成，不能作为成功交付。请查看本次命令的错误代码。\n')
        raise


def authentication(args):
    # Only called in the explicit retry/check stage, after browser login was checked.
    if args.cookie_stdin:
        jar = supplied_cookie_jar(sys.stdin.read())
        return jar, {'source': 'user-supplied-cookie', 'cookie_count': len(jar)}
    return browser_cookie_jar(args.browser, args.profile)


def request_note(url, jar=None):
    data, final, status, _ = fetch(url, jar)
    resolved = clean_url(final)
    try:
        note_id(url)
        target_url = url  # Keep the requested ID even if the response is a home/login shell.
    except ExtractError:
        try:
            note_id(resolved)
            target_url = resolved
        except ExtractError:
            raise ExtractError('NOTE_NOT_FOUND', '分享链接未返回可识别的目标笔记，请保留完整分享参数。') from None
    note = parse_page(data.decode('utf-8'), target_url)
    note['page_http_status'] = status
    return note


def needs_authentication(error):
    return error.code in ('NOTE_NOT_FOUND', 'INCOMPLETE_NOTE') or (
        error.code == 'HTTP_ERROR' and error.http_status in (401, 403)
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--url', action='append', default=[])
    parser.add_argument('--out', type=Path, default=Path('output/笔记提取'))
    parser.add_argument('--browser', choices=['chrome', 'edge', 'firefox', 'brave'], default='chrome')
    parser.add_argument('--profile')
    parser.add_argument('--authenticated', action='store_true', help='匿名失败且已检查浏览器登录后，进行Cookie认证重试')
    parser.add_argument('--cookie-stdin', action='store_true', help='认证重试或认证检查时，从stdin读取本次Cookie')
    parser.add_argument('--check-auth', action='store_true')
    parser.add_argument('--html', type=Path, help='离线复验，不证明 HTTP 抓取成功')
    args = parser.parse_args()
    if not args.url and not args.check_auth:
        parser.error('至少提供一个 --url')
    if args.html and (len(args.url) != 1 or args.check_auth or args.authenticated or args.cookie_stdin):
        parser.error('--html 必须和一个 --url 一起使用，不能同时使用认证参数')
    if args.cookie_stdin and not (args.authenticated or args.check_auth):
        parser.error('--cookie-stdin 仅用于 --authenticated 重试或 --check-auth 检查')
    if args.check_auth:
        try:
            jar, auth = authentication(args)
        except ExtractError as error:
            emit(status='auth_unavailable', code=error.code, message=str(error))
            return 2
        emit(status='auth_ready', **auth)
        return 0
    cached_auth = None
    cached_auth_error = None
    failed = auth_required = 0
    for input_index, raw in enumerate(args.url, 1):
        started = time.monotonic()
        try:
            url = clean_url(raw)
            jar, auth = None, {'source': 'anonymous', 'cookie_count': 0}
            if args.html:
                note = parse_page(args.html.read_text(encoding='utf-8'), url)
                note.update(page_http_status=None, extraction_method='offline-' + note['extraction_method'])
                auth = {'source': 'offline-html', 'cookie_count': 0}
            else:
                if args.authenticated:
                    if cached_auth is None and cached_auth_error is None:
                        try:
                            cached_auth = authentication(args)
                            emit(status='auth_ready', **cached_auth[1])
                        except ExtractError as error:
                            cached_auth_error = error
                    if cached_auth_error is not None:
                        failed += 1
                        auth_required += 1
                        emit(status='auth_unavailable', input_index=input_index,
                             code=cached_auth_error.code, message=str(cached_auth_error))
                        continue
                    jar, auth = cached_auth
                try:
                    note = request_note(url, jar)
                except ExtractError as error:
                    if not args.authenticated and needs_authentication(error):
                        failed += 1
                        auth_required += 1
                        emit(status='auth_required', input_index=input_index, code='AUTH_FALLBACK_REQUIRED', cause=error.code,
                             message='无Cookie请求未取得完整目标笔记。请在用户自己的浏览器检查登录，已登录后用--authenticated重试。')
                        continue
                    raise
            note['authentication'] = auth
            # Refresh only expired CDN failures, once, while retaining target ID.
            try:
                directory = export_note(note, args.out.resolve())
            except ExtractError as error:
                if error.code != 'HTTP_ERROR' or args.html:
                    raise
                emit(status='refresh_image_urls', note_id=note['note_id'])
                note = request_note(url, jar)
                note['authentication'] = auth
                directory = export_note(note, args.out.resolve())
            emit(status='exported', input_index=input_index, title=note['title'], note_id=note['note_id'], images=len(note['images']), directory=str(directory), seconds=round(time.monotonic() - started, 2))
        except ExtractError as error:
            failed += 1
            emit(status='failed', input_index=input_index, code=error.code, message=str(error))
        except Exception as error:
            failed += 1
            emit(status='failed', input_index=input_index, code='UNEXPECTED_ERROR', error_type=type(error).__name__, message='导出失败，请检查依赖与输出目录；未记录请求 Cookie。')
    return 2 if auth_required else 1 if failed else 0


if __name__ == '__main__':
    try:
        raise SystemExit(main())
    except ExtractError as error:
        emit(status='failed', code=error.code, message=str(error))
        raise SystemExit(2)
